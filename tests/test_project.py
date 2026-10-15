"""Run with: python -m unittest discover -s tests -v

No API key, network connection or extra testing library is needed.
"""

from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
import json
import re
import unittest
from unittest.mock import patch, MagicMock

import httpx
from docx import Document
from openai import OpenAI, OpenAIError
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
from werkzeug.datastructures import FileStorage

from ai_review import Review, get_ai_review, validate_evidence
from analyzer import analyze_resume, compare_job, extract_skills, has_metric, language_issues
from app import app
from parser import MAX_FILE_BYTES, parse_upload
from reports import report_as_text

ROOT = Path(__file__).resolve().parents[1]
RESUME = (ROOT / "samples/sample_resume.txt").read_text(encoding="utf-8")
JOB = (ROOT / "samples/sample_job.txt").read_text(encoding="utf-8")


def upload(data, filename):
    return FileStorage(stream=BytesIO(data), filename=filename)


def make_pdf(text="A fictional Python developer built a Flask application and tested the software with SQL data.", pages=1, password=None):
    """Create a tiny real PDF fixture using our existing pypdf dependency."""
    writer = PdfWriter()
    for number in range(pages):
        page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                 NameObject("/Subtype"): NameObject("/Type1"),
                                 NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        content = DecodedStreamObject()
        content.set_data(f"BT /F1 12 Tf 40 740 Td ({text}) Tj ET".encode("ascii"))
        page[NameObject("/Contents")] = writer._add_object(content)
    if password:
        writer.encrypt(password)
    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


class AnalysisTests(unittest.TestCase):
    def test_aliases_and_boundaries(self):
        found = extract_skills("JavaScript C++ C# natural language processing and Amazon Web Services")
        self.assertIn("JavaScript", found)
        self.assertNotIn("Java", found)
        self.assertIn("C++", found)
        self.assertIn("C#", found)
        self.assertIn("NLP", found)
        self.assertIn("AWS", found)

    def test_metrics_exclude_dates(self):
        self.assertFalse(has_metric("Developed software in 2024"))
        self.assertTrue(has_metric("Reduced runtime by 20%"))
        self.assertTrue(has_metric("Created 18 tests for the application"))

    def test_no_job_does_not_penalize_resume(self):
        report = analyze_resume(RESUME)
        self.assertIsNone(report["job"]["score"])
        self.assertIsNone(report["categories"][-1]["score"])
        self.assertIn("75 available", report["scoring_note"])
        expected = round(sum(c["score"] * c["weight"] for c in report["categories"][:-1]) / 75)
        self.assertEqual(report["overall_score"], expected)

    def test_formula_and_missing_skills(self):
        job = compare_job("Python Flask", "Python Flask Docker")
        self.assertEqual(job["missing"], ["Docker"])
        self.assertEqual(job["coverage"], 66.7)
        self.assertAlmostEqual(job["score"], round(0.75 * 100 * 2 / 3 + 0.25 * job["similarity"]), delta=1)

    def test_unknown_job_uses_similarity(self):
        job = compare_job("plant trees", "plant trees")
        self.assertIsNone(job["coverage"])
        self.assertEqual(job["score"], 100)

    def test_stopwords_do_not_crash(self):
        job = compare_job("the and or", "a the and")
        self.assertEqual(job["score"], 0)

    def test_repetition_does_not_inflate_skill_coverage(self):
        self.assertEqual(compare_job("Python " * 100, "Python SQL")["coverage"], 50)

    def test_structure_needs_headings(self):
        report = analyze_resume("I have education and skills and experience in programming.")
        self.assertEqual(report["sections"], [])

    def test_strong_sample_beats_weak(self):
        weak = (ROOT / "samples/weak_resume.txt").read_text(encoding="utf-8")
        self.assertGreater(analyze_resume(RESUME, JOB)["overall_score"], analyze_resume(weak, JOB)["overall_score"])

    def test_language_findings_are_specific(self):
        issues = language_issues("I have experiance with the the project.")
        self.assertTrue(any(i["text"] == "experiance" for i in issues))
        self.assertTrue(any(i["text"] == "the the" for i in issues))

    def test_scores_are_bounded_and_export_contains_checks(self):
        report = analyze_resume(RESUME, JOB)
        for category in report["categories"]:
            self.assertTrue(0 <= category["score"] <= 100)
        self.assertIn("Measurable outcomes", report_as_text(report))
        json.dumps(report)


class ParserTests(unittest.TestCase):
    def test_txt_bom(self):
        text, metadata = parse_upload(upload(RESUME.encode("utf-8-sig"), "RESUME.TXT"))
        self.assertTrue(text.startswith("ALEX"))
        self.assertEqual(metadata["source"], "TXT")

    def test_docx_preserves_body_order_and_warns_about_tables(self):
        document = Document()
        document.add_paragraph("Before the table: " + RESUME[:100])
        document.add_table(rows=1, cols=1).cell(0, 0).text = "Inside the table"
        document.add_paragraph("After the table")
        buffer = BytesIO()
        document.save(buffer)
        text, metadata = parse_upload(upload(buffer.getvalue(), "resume.docx"))
        self.assertLess(text.index("Before"), text.index("Inside"))
        self.assertLess(text.index("Inside"), text.index("After"))
        self.assertTrue(any("Tables" in message for message in metadata["warnings"]))

    def test_real_pdf(self):
        text, metadata = parse_upload(upload(make_pdf(), "resume.pdf"))
        self.assertIn("Python", text)
        self.assertEqual(metadata["pages"], 1)

    def test_encrypted_and_long_pdf(self):
        for data in [make_pdf(password="test"), make_pdf(pages=11)]:
            with self.assertRaises(ValueError):
                parse_upload(upload(data, "resume.pdf"))

    def test_blank_pdf_explains_ocr(self):
        with self.assertRaisesRegex(ValueError, "OCR"):
            parse_upload(upload(make_pdf(text=""), "scan.pdf"))

    def test_invalid_files_and_size(self):
        for data, name in [(b"bad", "bad.pdf"), (b"bad", "bad.docx"), (b"bad", "bad.exe"),
                           (b"", "empty.txt"), (b"\xff", "bad.txt"),
                           (b"a" * (MAX_FILE_BYTES + 1), "huge.txt"),
                           (("a" * 30001).encode(), "long.txt")]:
            with self.subTest(name=name), self.assertRaises(ValueError):
                parse_upload(upload(data, name))


class WebTests(unittest.TestCase):
    def setUp(self):
        app.config["TESTING"] = True
        self.client = app.test_client()
        self.client.get("/")
        with self.client.session_transaction() as session:
            self.token = session["csrf_token"]

    def post(self, **fields):
        return self.client.post("/analyze", data={"csrf_token": self.token, **fields})

    def test_home_and_demo(self):
        self.assertEqual(self.client.get("/").status_code, 200)
        self.assertIn(b"ALEX MORGAN", self.client.get("/?demo=1").data)

    def test_success_and_exports(self):
        response = self.post(resume_text=RESUME, job_description=JOB)
        self.assertEqual(response.status_code, 200)
        html = response.get_data(as_text=True)
        match = re.search(r'id="report-data">(.*?)</script>', html, re.S)
        self.assertIsNotNone(match)
        report = json.loads(match.group(1))
        self.assertEqual(report["overall_score"], analyze_resume(RESUME, JOB)["overall_score"])
        self.assertEqual(response.headers["Cache-Control"], "no-store")

    def test_upload_route(self):
        response = self.post(resume_file=(BytesIO(RESUME.encode()), "resume.txt"))
        self.assertEqual(response.status_code, 200)

    def test_validation_and_csrf(self):
        self.assertEqual(self.post().status_code, 400)
        self.assertEqual(self.post(resume_text="too short").status_code, 400)
        self.assertEqual(self.post(resume_text=RESUME, job_description="a" * 15001).status_code, 400)
        self.assertEqual(self.post(resume_text=RESUME, resume_file=(BytesIO(b"test"), "resume.txt")).status_code, 400)
        self.assertEqual(self.client.post("/analyze", data={"resume_text": RESUME}).status_code, 400)

    def test_request_limit(self):
        with patch.dict(app.config, {"MAX_CONTENT_LENGTH": 100}):
            response = self.post(resume_text="a" * 200)
        self.assertEqual(response.status_code, 413)

    def test_html_input_is_escaped(self):
        html = self.post(resume_text=RESUME + '\n<script>alert("test")</script>').get_data(as_text=True)
        self.assertNotIn('<script>alert("test")</script>', html)
        self.assertIn("&lt;script&gt;", html)

    @patch("app.get_ai_review")
    def test_ai_only_runs_with_opt_in(self, mocked):
        self.post(resume_text=RESUME)
        mocked.assert_not_called()
        mocked.return_value = {"status": "unavailable", "message": "Test fallback", "review": None}
        response = self.post(resume_text=RESUME, use_ai="yes")
        mocked.assert_called_once()
        self.assertIn(b"Test fallback", response.data)


class AITests(unittest.TestCase):
    def review(self):
        return Review(summary="Test summary", skill_checks=[], writing_suggestions=[],
                      bullet_rewrites=[], recommendations=["Use real evidence."])

    @patch.dict("os.environ", {"OPENAI_API_KEY": ""})
    def test_no_key(self):
        self.assertEqual(get_ai_review(RESUME, JOB)["status"], "unavailable")

    def test_evidence_filter_and_invented_numbers(self):
        review = Review(summary="Test", skill_checks=[
            {"skill": "Python", "status": "supported", "job_evidence": "Python", "resume_evidence": "Python", "explanation": "Mentioned"},
            {"skill": "Docker", "status": "supported", "job_evidence": "Docker", "resume_evidence": "Deployed with Docker", "explanation": "Invalid quote"},
        ], writing_suggestions=[], bullet_rewrites=[
            {"original": "Python", "suggestion": "Python for 5000 users", "reason": "Invented number"}
        ], recommendations=[])
        data = validate_evidence(review, RESUME, JOB)
        self.assertEqual(len(data["skill_checks"]), 1)
        self.assertEqual(data["bullet_rewrites"], [])
        self.assertEqual(data["discarded_items"], 2)

    @patch.dict("os.environ", {"OPENAI_API_KEY": "fake-test-key"})
    @patch("ai_review.OpenAI")
    def test_success_incomplete_and_provider_error(self, mocked):
        client = MagicMock()
        mocked.return_value.__enter__.return_value = client
        client.responses.parse.return_value = SimpleNamespace(status="completed", output_parsed=self.review())
        self.assertEqual(get_ai_review(RESUME, JOB)["status"], "complete")
        self.assertFalse(client.responses.parse.call_args.kwargs["store"])
        client.responses.parse.return_value = SimpleNamespace(status="incomplete", output_parsed=None)
        self.assertEqual(get_ai_review(RESUME, JOB)["status"], "unavailable")
        client.responses.parse.side_effect = OpenAIError("Test provider failure")
        self.assertEqual(get_ai_review(RESUME, JOB)["status"], "unavailable")

    @patch.dict("os.environ", {"OPENAI_API_KEY": "fake-test-key"})
    def test_real_sdk_with_fake_http_response(self):
        # Exercise the actual SDK's request schema and response parsing without a network call.
        sent = []

        def respond(request):
            sent.append(json.loads(request.content))
            return httpx.Response(200, json={
                "id": "resp_test", "object": "response", "created_at": 1,
                "status": "completed", "model": "gpt-4o-mini",
                "output": [{"id": "msg_test", "type": "message", "role": "assistant",
                            "status": "completed", "content": [{"type": "output_text",
                            "text": self.review().model_dump_json(), "annotations": []}]}],
            })

        client = OpenAI(api_key="fake-test-key", http_client=httpx.Client(transport=httpx.MockTransport(respond)))
        with patch("ai_review.OpenAI", return_value=client):
            result = get_ai_review(RESUME, JOB)
        self.assertEqual(result["status"], "complete")
        self.assertEqual(sent[0]["text"]["format"]["type"], "json_schema")
        self.assertFalse(sent[0]["store"])


if __name__ == "__main__":
    unittest.main()
